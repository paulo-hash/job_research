#!/usr/bin/env python3
"""Convert a cv_demo.md-shaped markdown file to a Word CV matching resume_example.docx."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x59, 0x59, 0x59)
RULE = "9A9A9A"
RIGHT_TAB = Twips(10282)
SKILL_TAB = Twips(1728)
PAGE_W = Twips(12240)
PAGE_H = Twips(15840)
MARGIN_TB = Twips(720)
MARGIN_LR = Twips(979)


def _set_run_font(run, *, size_pt: float, bold: bool = False, color: RGBColor = INK, spacing: int | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "Calibri")
    if spacing is not None:
        sp = rpr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            rpr.append(sp)
        sp.set(qn("w:val"), str(spacing))


def _bottom_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), RULE)
    pBdr.append(bottom)
    ppr.append(pBdr)


def _space(paragraph, *, before: int = 0, after: int = 0) -> None:
    paragraph.paragraph_format.space_before = Twips(before)
    paragraph.paragraph_format.space_after = Twips(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 1.25


def _add_tab_stop(paragraph, pos, alignment) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(pos, alignment)


def _split_bold(text: str) -> list[tuple[str, bool]]:
    parts: list[tuple[str, bool]] = []
    for i, chunk in enumerate(re.split(r"(\*\*.+?\*\*)", text)):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            parts.append((chunk[2:-2], True))
        else:
            parts.append((chunk, False))
    return parts or [("", False)]


def parse_cv_md(text: str) -> dict:
    lines = [line.rstrip() for line in text.splitlines()]
    name = ""
    header: list[str] = []
    tagline = ""
    sections: dict[str, list[str]] = {}
    current = ""
    body: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            name = line[2:].strip()
        elif line.startswith("## "):
            if current:
                sections[current] = body
            current = line[3:].strip().upper()
            body = []
        elif not current:
            stripped = line.strip()
            if stripped.startswith("**") and stripped.endswith("**"):
                tagline = stripped.strip("*").strip()
            elif stripped:
                header.append(stripped)
        else:
            body.append(line)
        i += 1
    if current:
        sections[current] = body
    return {"name": name, "header": header, "tagline": tagline, "sections": sections}


def _flush_role(doc, title: str, meta: str, bullets: list[str]) -> None:
    p = doc.add_paragraph()
    _space(p, before=110, after=10)
    _add_tab_stop(p, RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    title_run = p.add_run(title)
    _set_run_font(title_run, size_pt=10, bold=True)
    p.add_run("\t")
    meta_run = p.add_run(meta)
    _set_run_font(meta_run, size_pt=9, color=MUTED)
    for bullet in bullets:
        bp = doc.add_paragraph()
        _space(bp, after=50)
        bp.paragraph_format.left_indent = Twips(245)
        bp.paragraph_format.first_line_indent = Twips(-245)
        dash = bp.add_run("– ")
        _set_run_font(dash, size_pt=9.5, color=MUTED)
        text = bullet.lstrip("- ").strip()
        if text.startswith("– "):
            text = text[2:]
        for chunk, bold in _split_bold(text):
            run = bp.add_run(chunk)
            _set_run_font(run, size_pt=9.5, bold=bold)


def _write_experience(doc, lines: list[str]) -> None:
    title = ""
    meta = ""
    bullets: list[str] = []
    for line in lines:
        if line.startswith("### "):
            if title:
                _flush_role(doc, title, meta, bullets)
            title = line[4:].strip()
            meta = ""
            bullets = []
        elif line.startswith("- ") or line.startswith("– "):
            bullets.append(line)
        elif line.strip() and title and not meta:
            meta = line.strip()
    if title:
        _flush_role(doc, title, meta, bullets)


def _write_skills(doc, lines: list[str]) -> None:
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        match = re.match(r"\*\*(.+?)\*\*\s*(.*)$", stripped)
        if not match:
            continue
        p = doc.add_paragraph()
        _space(p, after=50)
        p.paragraph_format.left_indent = SKILL_TAB
        p.paragraph_format.first_line_indent = -SKILL_TAB
        _add_tab_stop(p, SKILL_TAB, WD_TAB_ALIGNMENT.LEFT)
        label = p.add_run(match.group(1).strip())
        _set_run_font(label, size_pt=9.5, bold=True)
        p.add_run("\t")
        value = p.add_run(match.group(2).strip())
        bold_value = match.group(1).strip().lower() == "languages"
        _set_run_font(value, size_pt=9.5, bold=bold_value)


def _write_summary(doc, lines: list[str]) -> None:
    text = " ".join(line.strip() for line in lines if line.strip())
    p = doc.add_paragraph()
    _space(p, after=40)
    for chunk, bold in _split_bold(text):
        run = p.add_run(chunk)
        _set_run_font(run, size_pt=9.5, bold=bold)


def _section_heading(doc, title: str) -> None:
    p = doc.add_paragraph()
    _space(p, before=130, after=70)
    _bottom_border(p)
    run = p.add_run(title)
    _set_run_font(run, size_pt=9, bold=True, spacing=28)


def build_docx(parsed: dict, dest: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.top_margin = MARGIN_TB
    section.bottom_margin = MARGIN_TB
    section.left_margin = MARGIN_LR
    section.right_margin = MARGIN_LR
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _space(name, after=40)
    name_run = name.add_run(parsed["name"])
    _set_run_font(name_run, size_pt=18, bold=True, spacing=32)

    for line in parsed["header"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _space(p, after=40)
        run = p.add_run(line)
        _set_run_font(run, size_pt=9.5)

    if parsed["tagline"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _space(p, after=140)
        run = p.add_run(parsed["tagline"])
        _set_run_font(run, size_pt=9.5, bold=True)

    writers = {
        "SUMMARY": _write_summary,
        "EXPERIENCE": _write_experience,
        "SKILLS": _write_skills,
        "EDUCATION": _write_experience,
    }
    for heading, lines in parsed["sections"].items():
        _section_heading(doc, heading)
        writer = writers.get(heading)
        if writer:
            writer(doc, lines)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown")
    parser.add_argument("docx")
    args = parser.parse_args(argv)
    src = Path(args.markdown)
    if not src.exists():
        print(f"missing {src}", file=sys.stderr)
        return 1
    parsed = parse_cv_md(src.read_text(encoding="utf-8"))
    if not parsed["name"]:
        print("markdown has no # name heading", file=sys.stderr)
        return 1
    build_docx(parsed, Path(args.docx))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
